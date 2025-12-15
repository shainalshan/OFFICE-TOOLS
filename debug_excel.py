import openpyxl
import os

def inspect_excel():
    filename = 'Asset list/PIXL DEVICES (2).xlsx'
    if not os.path.exists(filename):
        print("File not found.")
        return

    wb = openpyxl.load_workbook(filename, data_only=True)
    ws = wb.active
    
    print("SHEET NAMES:", wb.sheetnames)
    
    target_sheets = ['DUBAI DEVICE LIST', 'INDIA DEVICE LIST']
    
    for sheet_name in target_sheets:
        if sheet_name in wb.sheetnames:
            print(f"\n--- INSPECTING {sheet_name} ---")
            ws = wb[sheet_name]
            for i, row in enumerate(ws.iter_rows(min_row=1, max_row=7, values_only=True)):
                print(f"\nROW {i+1}:")
                for idx, val in enumerate(row[:10]):
                    if val:
                        print(f"  [{idx}] {val}")

if __name__ == "__main__":
    inspect_excel()
