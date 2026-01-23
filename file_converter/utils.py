import os
import threading
import uuid
from django.core.cache import cache
from django.conf import settings
import io

# Safe top-level imports where likely no system deps needed
try:
    from PIL import Image
except ImportError:
    Image = None

try:
    import pandas as pd
except ImportError:
    pd = None

try:
    from pypdf import PdfReader, PdfWriter
except ImportError:
    PdfReader = None
    PdfWriter = None

class FileConverter:
    def __init__(self, file_path, task_id):
        self.file_path = file_path
        self.task_id = task_id
        self.filename = os.path.basename(file_path)

    def update_status(self, message, status="processing"):
        cache.set(f"task_{self.task_id}", {"status": status, "message": message}, timeout=360)

    def run_pdf_to_word(self):
        try:
            if not self.file_path.lower().endswith('.pdf'):
                 self.update_status("Error: Input file must be a PDF.", status="failed")
                 return None

            from pdf2docx import Converter
            self.update_status("Starting PDF to Word conversion...")
            docx_path = self.file_path.replace('.pdf', '.docx')
            cv = Converter(self.file_path)
            cv.convert(docx_path, start=0, end=None)
            cv.close()
            self.update_status("Conversion complete.", status="completed")
            return docx_path
        except ImportError:
             self.update_status("Error: pdf2docx library not found.", status="failed")
        except Exception as e:
            self.update_status(f"Error: {str(e)}", status="failed")

    def run_pdf_to_pptx(self):
        try:
            if not self.file_path.lower().endswith('.pdf'):
                 self.update_status("Error: Input file must be a PDF.", status="failed")
                 return None

            from pdf2image import convert_from_path
            from pptx import Presentation
            from pptx.util import Inches

            self.update_status("Converting PDF to images...")
            # Note: poppler must be installed on the system for this to work
            try:
                images = convert_from_path(self.file_path)
            except Exception as e:
                import subprocess
                if "Unable to get page count" in str(e) or "poppler" in str(e).lower():
                     raise ImportError(f"Poppler is likely not installed or not in PATH. Please install Poppler for Windows and add it to your PATH environment variable. \nOriginal Error: {e}")
                raise e
            prs = Presentation()
            
            total_pages = len(images)
            for i, image in enumerate(images):
                self.update_status(f"Processing page {i+1} of {total_pages}...")
                slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide
                image_path = f"{self.file_path}_temp_{i}.jpg"
                image.save(image_path, 'JPEG')
                
                left = top = Inches(0)
                slide.shapes.add_picture(image_path, left, top, height=Inches(7.5))
                try:
                    os.remove(image_path)
                except:
                    pass
                
            pptx_path = self.file_path.replace('.pdf', '.pptx')
            prs.save(pptx_path)
            self.update_status("Conversion complete.", status="completed")
            return pptx_path
        except ImportError as e:
            self.update_status(f"Error: Library missing ({str(e)})", status="failed")
        except Exception as e:
            self.update_status(f"Error: {str(e)}", status="failed")

    def run_excel_to_pdf(self):
        try:
            if not self.file_path.lower().endswith(('.xlsx', '.xls')):
                 self.update_status("Error: Input file must be an Excel file (.xlsx, .xls).", status="failed")
                 return None

            if pd is None:
                raise ImportError("pandas not installed")
            
            # Lazy import weasyprint as it often causes system dep issues (GTK3)
            from weasyprint import HTML

            self.update_status("Reading Excel file...")
            if self.file_path.endswith('.xlsx'):
                df = pd.read_excel(self.file_path, engine='openpyxl')
            else:
                df = pd.read_excel(self.file_path)
                
            self.update_status("Generating PDF...")
            html_string = df.to_html()
            pdf_path = self.file_path + '.pdf'
            HTML(string=html_string).write_pdf(pdf_path)
            self.update_status("Conversion complete.", status="completed")
            return pdf_path
        except ImportError as e:
            self.update_status(f"Error: Library missing ({str(e)}) - WeasyPrint requires GTK3 runtime on Windows.", status="failed")
        except Exception as e:
            self.update_status(f"Error: {str(e)}", status="failed")

    def run_pdf_to_excel(self):
        try:
            if not self.file_path.lower().endswith('.pdf'):
                 self.update_status("Error: Input file must be a PDF.", status="failed")
                 return None

            import pdfplumber
            if pd is None:
                raise ImportError("pandas not installed")

            self.update_status("Extracting tables from PDF...")
            final_df = pd.DataFrame()
            with pdfplumber.open(self.file_path) as pdf:
                total_pages = len(pdf.pages)
                for i, page in enumerate(pdf.pages):
                    self.update_status(f"Processing page {i+1} of {total_pages}...")
                    table = page.extract_table()
                    if table:
                        df = pd.DataFrame(table[1:], columns=table[0])
                        final_df = pd.concat([final_df, df])
            
            excel_path = self.file_path.replace('.pdf', '.xlsx')
            final_df.to_excel(excel_path, index=False)
            self.update_status("Conversion complete.", status="completed")
            return excel_path
        except ImportError as e:
             self.update_status(f"Error: Library missing ({str(e)})", status="failed")
        except Exception as e:
            self.update_status(f"Error: {str(e)}", status="failed")


    def run_pptx_to_pdf(self):
        try:
            if not self.file_path.lower().endswith(('.pptx', '.ppt')):
                 self.update_status("Error: Input file must be a PowerPoint presentation (.pptx, .ppt).", status="failed")
                 return None

            import comtypes.client
            import pythoncom
            
            # Initialize COM for this thread
            pythoncom.CoInitialize()

            self.update_status("Opening PowerPoint...")
            powerpoint = comtypes.client.CreateObject("Powerpoint.Application")
            powerpoint.Visible = 1
            
            self.update_status("Converting Presentation...")
            # Must use absolute path for COM
            abs_path = os.path.abspath(self.file_path)
            output_path = abs_path.replace('.pptx', '.pdf').replace('.ppt', '.pdf')
            
            deck = powerpoint.Presentations.Open(abs_path)
            # Format 32 is ppSaveAsPDF
            deck.SaveAs(output_path, 32)
            deck.Close()
            powerpoint.Quit()
            
            pythoncom.CoUninitialize()
            
            self.update_status("Conversion complete.", status="completed")
            return output_path
            
        except ImportError as e:
            if 'comtypes' in str(e):
                 self.update_status("Error: 'comtypes' library not found. Please install it.", status="failed")
            elif 'pythoncom' in str(e) or 'pywin32' in str(e):
                 self.update_status("Error: 'pywin32' library not found. Please install it.", status="failed")
            else:
                 self.update_status(f"Error: Library missing ({str(e)})", status="failed")
        except Exception as e:
            self.update_status(f"Error: {str(e)} (Ensure MS PowerPoint is installed)", status="failed")

    def run_compression(self, file_type):
        try:
            self.update_status("Starting compression...")
            output_path = ""
            if file_type == 'pdf':
                if PdfReader is None:
                     raise ImportError("pypdf not installed")

                reader = PdfReader(self.file_path)
                writer = PdfWriter()
                
                for page in reader.pages:
                    # Compress content streams
                    page.compress_content_streams()
                    writer.add_page(page)
                
                writer.compress_identical_objects()
                
                output_path = self.file_path.replace('.pdf', '_compressed.pdf')
                
                with open(output_path, "wb") as f:
                    writer.write(f)
                    
            elif file_type == 'image':
                if Image is None:
                    raise ImportError("Pillow not installed")
                
                img = Image.open(self.file_path)
                output_path = f"compressed_{self.filename}"
                output_path = os.path.join(os.path.dirname(self.file_path), output_path)
                img.save(output_path, optimize=True, quality=50)
                
            self.update_status("Compression complete.", status="completed")
            return output_path
        except ImportError as e:
             self.update_status(f"Error: Library missing ({str(e)})", status="failed")
        except Exception as e:
            self.update_status(f"Error: {str(e)}", status="failed")

def handle_file_operation(operation, file_path, task_id):
    converter = FileConverter(file_path, task_id)
    result_path = None
    
    try:
        if operation == 'pdf_to_docx':
            result_path = converter.run_pdf_to_word()
        elif operation == 'pdf_to_pptx':
            result_path = converter.run_pdf_to_pptx()
        elif operation == 'pptx_to_pdf':
            result_path = converter.run_pptx_to_pdf()
        elif operation == 'excel_to_pdf':
            result_path = converter.run_excel_to_pdf()
        elif operation == 'pdf_to_excel':
            result_path = converter.run_pdf_to_excel()
        elif operation == 'compress_pdf':
            result_path = converter.run_compression('pdf')
        elif operation == 'compress_image':
            result_path = converter.run_compression('image')
            
        if result_path and os.path.exists(result_path):
             cache.set(f"result_{task_id}", result_path, timeout=3600)
             
    except Exception as e:
        converter.update_status(f"Fatal error: {str(e)}", status="failed")
