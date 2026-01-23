import os
import django
import sys
import openpyxl
from datetime import datetime

# Setup Django Environment
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
django.setup()

from assets.models import Asset
from django.contrib.auth.models import User

def import_assets_from_excel(filename='Asset list/PIXL DEVICES (2).xlsx'):
    file_path = os.path.join(os.getcwd(), filename)
    if not os.path.exists(file_path):
        print(f"Error: File {filename} not found.")
        return

    print(f"Reading {filename}...")
    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    target_sheets = ['DUBAI DEVICE LIST', 'INDIA DEVICE LIST', 'PAKISTAN DEVICE LIST']
    
    created_count = 0
    updated_count = 0
    
    for sheet_name in target_sheets:
        if sheet_name not in wb.sheetnames:
            print(f"Skipping {sheet_name} (Not found)")
            continue
            
        ws = wb[sheet_name]
        print(f"\n--- PROCESSING {sheet_name} ---")
        
        # Determine Mapping based on sheet
        if sheet_name == 'DUBAI DEVICE LIST':
            # Dubai: [0]ID, [1]Staff, [2]Type, [3]Brand, [4]Model, [5]Serial, [6]Signed?
            idx_staff = 1
            idx_type = 2
            idx_brand = 3
            idx_model = 4
            idx_serial = 5
            idx_mni = 0
            idx_remarks = 9 # Status/Remarks
        else:
            # India/Pakistan: [0]ID, [1]Type, [2]Brand, [3]Model, [4]Serial, [5]Staff, [6]Status
            idx_staff = 5
            idx_type = 1
            idx_brand = 2
            idx_model = 3
            idx_serial = 4
            idx_mni = 0
            idx_remarks = 9 # Assuming same column index for "Good working condition" seen in debug

        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
            # Ensure row length
            if not row or len(row) < 6:
                continue

            try:
                serial = str(row[idx_serial]).strip() if row[idx_serial] else ""
                if not serial or serial.lower() in ['none', 'n/a', 'serial number']:
                    continue
                    
                staff_name = str(row[idx_staff]).strip() if row[idx_staff] else ""
                device_type_raw = str(row[idx_type]).strip() if row[idx_type] else ""
                brand = str(row[idx_brand]).strip() if row[idx_brand] else "Generic"
                model_detail = str(row[idx_model]).strip() if row[idx_model] else ""
                mni = str(row[idx_mni]).strip() if row[idx_mni] else ""
                remarks = str(row[idx_remarks]).strip() if row[idx_remarks] else ""
                
                # Normalize Data
                device_type = 'OTHER'
                os_type = 'NONE'
                dt_lower = device_type_raw.lower() if device_type_raw else ""
                detail_lower = model_detail.lower() if model_detail else ""
                
                if 'laptop' in dt_lower:
                    if 'apple' in brand.lower():
                        device_type = 'MAC'
                        os_type = 'MAC'
                    else:
                        device_type = 'WINDOWS'
                        os_type = 'WINDOWS'
                elif 'monitor' in dt_lower or 'screen' in dt_lower:
                    device_type = 'MONITOR' if 'MONITOR' in [c[0] for c in Asset.DEVICE_TYPES] else 'OTHER'
                elif 'mouse' in dt_lower or 'keyboard' in dt_lower:
                    device_type = 'ACCESSORY' if 'ACCESSORY' in [c[0] for c in Asset.DEVICE_TYPES] else 'OTHER'
                elif 'iphone' in detail_lower or ('apple' in detail_lower and ('gb' in detail_lower or '16' in detail_lower)):
                    brand = "Apple"
                    device_type = 'IPHONE'
                    os_type = 'MAC'

                # Truncate fields to model limits
                brand = brand[:50]
                model_detail = model_detail[:100]
                mni = mni[:50]
                remarks = remarks # No limit on TextField usually, but good to know
                
                # Status Logic
                status = 'IN_STORE'
                if staff_name and staff_name.lower() not in ['in store', 'stock', 'office']:
                    status = 'IN_USE'
                
                # Find User match
                user = None
                if status == 'IN_USE':
                    # Simple first name match
                    fname = staff_name.split(' ')[0]
                    user = User.objects.filter(first_name__icontains=fname).first()

                # Location
                location = 'DUBAI'
                if 'INDIA' in sheet_name: location = 'INDIA'
                elif 'PAKISTAN' in sheet_name: location = 'PAKISTAN'

                asset, created = Asset.objects.get_or_create(
                    serial_number=serial,
                    defaults={
                        'mni': mni,
                        'brand': brand,
                        'model_detail': model_detail,
                        'device_type': device_type,
                        'os_type': os_type,
                        'status': status,
                        'assigned_to': user,
                        'assigned_to_email': staff_name,
                        'location': location
                    }
                )

                if created:
                    created_count += 1
                    print(f"[CREATED] {serial} ({staff_name})")
                else:
                    # Update existing
                    asset.device_type = device_type
                    asset.brand = brand
                    asset.model_detail = model_detail
                    asset.os_type = os_type
                    asset.location = location
                    asset.status = status
                    asset.assigned_to = user
                    asset.assigned_to_email = staff_name
                    asset.save()
                    updated_count += 1
                    # print(f"[UPDATED] {serial}")
            
            except Exception as e:
                print(f"Error Row {i}: {e}")

    print(f"\n--- TOTAL SUMMARY ---")
    print(f"Created: {created_count}")
    print(f"Updated: {updated_count}")

if __name__ == '__main__':
    import_assets_from_excel()
