import openpyxl
import os

def inspect_staff():
    filename = 'Asset list/PIXL DEVICES (2).xlsx'
    if not os.path.exists(filename):
        print("File not found.")
        return

    wb = openpyxl.load_workbook(filename, data_only=True)
    if 'STAFF LIST' not in wb.sheetnames:
        print("STAFF LIST sheet not found!")
        return
        
    ws = wb['STAFF LIST']
    
    print("\n--- STAFF LIST INSPECTION ---")
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True)):
        print(f"Row {i+1}: {row}")

if __name__ == "__main__":
    inspect_staff()
