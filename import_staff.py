import os
import django
import sys
import openpyxl

# Setup Django Environment
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
django.setup()

from contacts.models import Contact

def import_staff_from_excel(filename='Asset list/PIXL DEVICES (2).xlsx'):
    file_path = os.path.join(os.getcwd(), filename)
    if not os.path.exists(file_path):
        print(f"Error: File {filename} not found.")
        return

    print(f"Reading {filename}...")
    wb = openpyxl.load_workbook(file_path, data_only=True)
    if 'STAFF LIST' not in wb.sheetnames:
        print("STAFF LIST sheet not found.")
        return
        
    ws = wb['STAFF LIST']
    
    # Based on inspection:
    # Row 1 seems to have headers but they are messy.
    # Data starts Row 2.
    # Column 1 looks like Name (Index 1).
    
    count = 0
    updated = 0
    
    # Inspecting finding showed names in Column index 1 (B)
    # Row 2: [0]Staff No, [1]Name, [2]Pixl Devices, [3]Personal Devices, [4]Total
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 2:
            continue
            
        staff_no = str(row[0]).strip() if row[0] else None
        name = str(row[1]).strip() if row[1] else ""
        personal_devices = str(row[3]).strip() if len(row) > 3 and row[3] else None
        
        # Skip header-like rows or empty names
        if not name or name.lower() in ['name', 'total', 'none'] or name.lower().startswith('dubai'):
            continue
            
        # Check against existing Contact by Name
        contact, created = Contact.objects.get_or_create(name__iexact=name, defaults={'name': name})
        
        # Update details
        contact.staff_no = staff_no
        contact.personal_devices = personal_devices
        if 'designation' not in contact.__dict__ or not contact.designation:
             contact.designation = 'Staff'
        contact.save()
        
        if created:
            count += 1
            print(f"[CREATED] {name} (No: {staff_no})")
        else:
            updated += 1
            # print(f"[UPDATED] {name}")
            
    print(f"\n--- IMPORT COMPLETE ---")
    print(f"Created: {count}")
    print(f"Updated: {updated}")

if __name__ == '__main__':
    import_staff_from_excel()
